#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""端到端自测：验证本技能的核心链路真的能跑通（渲染 → 回读校验），并做反向用例。

覆盖：
    1 环境自检可运行并给出推荐解释器
    2 环境状态机四态判定正确（ready / partial / needs_setup / unavailable）
    3 docx_helpers 首行缩进与节点删除生效
    4 标题恒在表格上方
    5 单元格先裁剪再写：模板提示文字被清除且末行不丢
    6 渲染课程标准 docx 并通过回读校验
    7 反向用例：排版被破坏时校验器必须报 FAIL
    8 Markdown 降级产出含版式对照表

用法：
    python selftest.py                 # 产物写到系统临时目录
    python selftest.py --out <目录>    # 产物落到指定目录，便于留证据
    python selftest.py --json

退出码：全部通过为 0，任一失败为 1。全程不启子进程、不访问网络、不删除文件。
"""
from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(SCRIPT_DIR))

import check_environment as env                                  # noqa: E402
import export_markdown as exporter                               # noqa: E402
import render_course_standard as renderer                        # noqa: E402
import validate_docx as validator                                # noqa: E402
from docx import Document                                        # noqa: E402
from docx.oxml.ns import qn                                      # noqa: E402
from docx_helpers import (                                       # noqa: E402
    drop_element, insert_heading_table_before, set_first_line_chars, write_cell,
)

RESULTS = []


def record(name, ok, detail):
    RESULTS.append({"case": name, "ok": bool(ok), "detail": detail})
    print(f"  [{'PASS' if ok else 'FAIL'}] {name} —— {detail}")


def case_environment(out_dir):
    report = env.build_report(SKILL_DIR)
    ok = report["status"] in {"ready", "partial", "needs_setup", "unavailable"}
    record("1 环境自检", ok,
           f"status={report['status']} 解释器={report['recommended_interpreter']}")
    (out_dir / "env-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return ok


def case_state_machine():
    ready = [{"required": True, "status": "ready"}, {"required": False, "status": "ready"}]
    partial = [{"required": True, "status": "ready"}, {"required": False, "status": "missing"}]
    needs = [{"required": True, "status": "missing"}]
    broken = [{"required": True, "status": "unavailable"}]
    got = [env.summarize(ready), env.summarize(partial), env.summarize(needs), env.summarize(broken)]
    want = ["ready", "partial", "needs_setup", "unavailable"]
    record("2 环境状态机四态", got == want, f"实测={got}")


def case_helpers(out_dir):
    doc = Document()
    p = doc.add_paragraph()
    set_first_line_chars(p, 2, size_pt=16)
    ind = p._p.find(qn('w:pPr')).find(qn('w:ind'))
    flc = ind.get(qn('w:firstLineChars')) if ind is not None else None

    holder = doc.add_paragraph("将被删除")
    holder_el = holder._p
    drop_element(holder_el)
    still_there = holder_el.getparent() is not None

    ok = flc == '200' and not still_there
    record("3 缩进与节点删除", ok, f"firstLineChars={flc} 删除后仍挂树={still_there}")
    return ok


def case_heading_above_table():
    doc = Document()
    insert_heading_table_before(doc, None, "表1 学时分配表",
                                [["序号", "学时"], ["1", "16"]])
    names = [child.tag.split('}')[-1] for child in doc.element.body]
    p_tag, tbl_tag = qn('w:p'), qn('w:tbl')
    order = [child.tag for child in doc.element.body]
    ok = (len(doc.tables) == 1 and order.index(p_tag) < order.index(tbl_tag))
    record("4 标题在表格上方", ok, f"body 顺序={names}，表格数={len(doc.tables)}")


def case_strip_cell_keeps_last_line():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].add_run("教学目的及要求：")
    for _ in range(3):
        cell.add_paragraph("（需包括但不限于以下内容：）")
    lines = [f"第{i}行内容" for i in range(1, 6)]
    write_cell(cell, lines)
    texts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    ok = (texts == lines) and not any("需包括但不限于" in t for t in texts)
    record("5 单元格先裁剪再写", ok, f"写入{len(lines)}行 实得{len(texts)}行 末行={texts[-1] if texts else '空'}")


def case_render_and_validate(out_dir):
    target = out_dir / "demo-course-standard.docx"
    doc = renderer.build_document(renderer.DEMO, SKILL_DIR / "templates" / "课程标准.docx")
    doc.save(str(target))
    report = validator.validate(target, "course-standard")
    (out_dir / "validate-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    record("6 渲染并回读校验", report["passed"],
           f"{target.name} 段落={report['paragraphs_checked']} 表格={report['tables']} 失败={report['failure_count']}")
    return target, report["passed"]


def case_negative(out_dir, source: Path):
    """反向用例：故意破坏排版，校验器必须报 FAIL，否则校验形同虚设。"""
    broken = out_dir / "broken-course-standard.docx"
    doc = Document(str(source))
    hit = 0
    for p in doc.paragraphs:
        if p.text.strip().startswith("（一）"):
            for run in p.runs:
                rPr = run._element.get_or_add_rPr()
                rPr.get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
            ind = p._p.get_or_add_pPr().find(qn('w:ind'))
            if ind is not None:
                ind.set(qn('w:firstLineChars'), '0')
            hit += 1
            break
    doc.save(str(broken))
    report = validator.validate(broken, "course-standard")
    ok = hit == 1 and not report["passed"] and report["failure_count"] >= 1
    detail = f"注入{hit}处破坏 → 校验报错{report['failure_count']}项"
    if report["failures"]:
        detail += f"；首条：{report['failures'][0]['expected']} / 实测 {report['failures'][0]['actual']}"
    record("7 反向用例能被抓住", ok, detail)


def case_markdown_fallback(out_dir):
    text = exporter.render(exporter.DEMO, "standard")
    target = out_dir / "demo-course-standard.md"
    target.write_text(text, encoding="utf-8")
    ok = "附：Word 版式对照表" in text and "| 元素 | 字体 |" in text
    record("8 Markdown 降级产出", ok, f"{target.name} 共 {len(text)} 字符，含版式对照表")


def main() -> int:
    ap = argparse.ArgumentParser(description="教学文件写作套件 · 端到端自测")
    ap.add_argument("--out", type=Path, help="产物输出目录，缺省用系统临时目录")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    out_dir = args.out or Path(tempfile.mkdtemp(prefix="suit-sz-selftest-"))
    out_dir.mkdir(parents=True, exist_ok=True)

    print("=" * 62)
    print(" 教学文件写作套件 · 端到端自测")
    print(f" 技能目录：{SKILL_DIR}")
    print(f" 产物目录：{out_dir}")
    print("=" * 62)

    case_environment(out_dir)
    case_state_machine()
    case_helpers(out_dir)
    case_heading_above_table()
    case_strip_cell_keeps_last_line()
    _target, render_ok = case_render_and_validate(out_dir)
    case_negative(out_dir, _target)
    case_markdown_fallback(out_dir)

    passed = sum(1 for item in RESULTS if item["ok"])
    summary = {"passed": passed, "total": len(RESULTS),
               "ok": passed == len(RESULTS) and render_ok,
               "out_dir": str(out_dir), "results": RESULTS}
    print("-" * 62)
    print(f" 结果：{passed}/{len(RESULTS)} 项通过")
    print("=" * 62)

    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if summary["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
