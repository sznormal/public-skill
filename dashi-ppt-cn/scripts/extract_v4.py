#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_v4.py — 从 dashi-ppt export:pptx 的 4N 页 deck 中提取干净的 N 页可编辑 PPTX

背景：
  dashi-ppt 在 variantOutputMode:"comparison" 导出时，每逻辑页输出 4 个变体
  （v1 模板 / v2 模板 / v3 模板 / v4 定制），共 4N 页。
  实测 theme09 的 v1-v3 模板变体常残留未填占位（text-1/text-2）与主题自带英文装饰
  文案（如 "ANNUAL RESEARCH REPORT" / "FEATURED COMPANIES"）；v4 定制版内容完整、
  无泄漏。故教材交付应取每组第 4 页（v4）。

用法：
  python extract_v4.py <in.pptx> <out.pptx>

依赖：python-pptx（managed venv：C:/Users/29206/.workbuddy/binaries/python/envs/default）
  "C:/Users/29206/.workbuddy/binaries/python/envs/default/Scripts/python.exe" extract_v4.py in.pptx out.pptx
"""
import sys
import re
from pptx import Presentation
from pptx.oxml.ns import qn

# 中英文泄漏 / 占位特征词（theme09 模板默认装饰与未填占位）。
# 仅匹配精确装饰短语，避免误伤教材真实内容（如 dev_env_report.txt 含 REPORT 子串）。
LEAK = re.compile(
    r'ANNUAL RESEARCH REPORT|FEATURED COMPANIES|CONTACT SHEET|SoundWave'
    r'|影像走带|声浪'
    r'|text-1|text-2|text-3|请输入文本|Key Metrics|Roadmap|End of Report',
    re.I,
)


def main():
    if len(sys.argv) < 3:
        print("用法: python extract_v4.py <in.pptx> <out.pptx>")
        sys.exit(2)

    src, out = sys.argv[1], sys.argv[2]
    prs = Presentation(src)
    n = len(prs.slides)
    if n % 4 != 0:
        print(f"[extract_v4] 警告: 输入 {n} 页不是 4 的倍数，可能不是 comparison 导出；"
              f"将仍按每组第4页提取。")

    # v4 是每组(4页)的第4个：索引 3, 7, 11, ..., 4k-1
    keep = set(range(3, n, 4))
    xml = prs.slides._sldIdLst
    sld_ids = list(xml)
    for idx in range(len(sld_ids) - 1, -1, -1):
        if idx not in keep:
            rId = sld_ids[idx].get(qn('r:id'))
            xml.remove(sld_ids[idx])
            prs.part.drop_rel(rId)

    prs.save(out)

    # 校验泄漏
    p2 = Presentation(out)
    leaks = 0
    titles = []
    for i, s in enumerate(p2.slides):
        first = ""
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                first = sh.text_frame.text.strip().split("\n")[0][:40]
                break
        titles.append(first)
        for sh in s.shapes:
            if sh.has_text_frame and LEAK.search(sh.text_frame.text):
                leaks += 1
                print(f"  [警告] 第{i+1}页疑似泄漏: {sh.text_frame.text.strip()[:60]}")
                break

    print(f"[extract_v4] 输入 {n} 页 -> 提取 {len(p2.slides)} 页，输出: {out}")
    if leaks:
        print(f"[extract_v4] 发现 {leaks} 处疑似英文/占位泄漏，请检查 v4 内容是否完整")
    else:
        print("[extract_v4] 未检测到英文泄漏/占位，提取干净")


if __name__ == "__main__":
    main()
